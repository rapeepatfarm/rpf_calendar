"""routers/history.py — ประวัติงานที่ผ่านมาและสรุปสถิติ + ส่งออก Excel"""
from datetime import date, timedelta
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Request
from fastapi.responses import RedirectResponse, StreamingResponse

from auth import require_login
from forms import OptInt
from database import fetchall, get_conn
from services import reports, thaidate
from view import page

router = APIRouter(prefix="/history")

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _range(preset: str, start: str, end: str) -> tuple[date, date, str]:
    """แปลงตัวเลือกช่วงเวลาเป็นวันเริ่ม-จบ · คืนชื่อช่วงไว้แสดงหัวรายงานด้วย"""
    today = thaidate.today()
    month_start = today.replace(day=1)

    if preset == "last_month":
        last_end = month_start - timedelta(days=1)
        return last_end.replace(day=1), last_end, "เดือนที่แล้ว"
    if preset == "last_90":
        return today - timedelta(days=89), today, "90 วันล่าสุด"
    if preset == "year":
        return date(today.year, 1, 1), date(today.year, 12, 31), f"ปี {today.year + 543}"
    if preset == "custom":
        a, b = thaidate.parse_iso(start), thaidate.parse_iso(end)
        if a and b and a <= b:
            return a, b, f"{thaidate.short(a)} – {thaidate.short(b)}"
        # ช่วงที่กรอกมาใช้ไม่ได้ ตกกลับไปเดือนนี้แทนการเด้ง error ใส่หน้าผู้ใช้
    return month_start, thaidate.add_months(month_start, 1) - timedelta(days=1), "เดือนนี้"


PRESETS = (("this_month", "เดือนนี้"), ("last_month", "เดือนที่แล้ว"),
           ("last_90", "90 วันล่าสุด"), ("year", "ทั้งปี"))


@router.get("")
def history_page(request: Request, preset: str = "this_month",
                 start: str = "", end: str = "",
                 assignee_id: OptInt = None, category_id: OptInt = None):
    user = require_login(request)
    a, b, label = _range(preset, start, end)
    filters = {"assignee_id": assignee_id, "category_id": category_id, "status": None}

    with get_conn() as conn:
        staff = fetchall(conn, "SELECT id, name FROM staff WHERE active ORDER BY sort_order, name")
        categories = fetchall(conn, "SELECT id, name, color FROM activity_categories "
                                    "WHERE active ORDER BY sort_order, name")
        rows = reports.rows_in_range(conn, a, b, filters)
        return page(request, user, "history.html", conn=conn,
                    summary=reports.summary(conn, a, b, filters),
                    people=reports.by_person(conn, a, b, filters),
                    cats=reports.by_category(conn, a, b, filters),
                    rows=rows, capped=len(rows) >= 500,
                    start=a, end=b, label=label, preset=preset,
                    presets=PRESETS, f=filters, staff=staff, categories=categories,
                    OUTCOME_LABELS=reports.OUTCOME_LABELS)


@router.get("/export")
def export_excel(request: Request, preset: str = "this_month",
                 start: str = "", end: str = "",
                 assignee_id: OptInt = None, category_id: OptInt = None):
    """ส่งออกเป็น Excel 3 ชีต: สรุป · รายคน · รายการทั้งหมด"""
    require_login(request)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return RedirectResponse(
            "/history?err=" + quote("ยังไม่ได้ติดตั้ง openpyxl — สั่ง pip install openpyxl"),
            status_code=303)

    a, b, label = _range(preset, start, end)
    filters = {"assignee_id": assignee_id, "category_id": category_id, "status": None}

    with get_conn() as conn:
        total = reports.summary(conn, a, b, filters)
        people = reports.by_person(conn, a, b, filters)
        # ส่งออกไม่จำกัด 500 แถวเหมือนบนหน้าจอ — คนโหลดไฟล์ต้องการข้อมูลครบ
        rows = reports.rows_in_range(conn, a, b, filters, limit=None)

    book = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F7DE1")

    def write_sheet(sheet, headers, data, widths):
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for record in data:
            sheet.append(record)
        for i, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = width
        sheet.freeze_panes = "A2"

    # ── ชีตสรุป ──
    s1 = book.active
    s1.title = "สรุป"
    pct = "-" if total["on_time_pct"] is None else f"{total['on_time_pct']}%"
    write_sheet(s1, ["หัวข้อ", "ค่า"], [
        ["ช่วงเวลา", f"{thaidate.short(a)} – {thaidate.short(b)} ({label})"],
        ["กิจกรรมทั้งหมด", total["total"]],
        ["เสร็จตรงแผน", total["on_time"]],
        ["เสร็จช้ากว่าแผน", total["late"]],
        ["ไม่ได้ทำ", total["missed"]],
        ["ยกเลิก", total["cancelled"]],
        ["ยังไม่ปิด", total["still_open"]],
        ["% ทำตรงแผน", pct],
        ["ช้าเฉลี่ย (วัน)", total["avg_late_days"]],
        ["เวลาที่ใช้จริงรวม (ชม.)", round(total["minutes"] / 60, 1)],
    ], [26, 34])
    s1["B8"].font = Font(bold=True)

    # ── ชีตรายคน ──
    s2 = book.create_sheet("รายคน")
    write_sheet(s2, ["ผู้รับผิดชอบ", "ตำแหน่ง", "ทั้งหมด", "ตรงแผน", "ช้า",
                     "ไม่ได้ทำ", "ยกเลิก", "ยังไม่ปิด", "% ตรงแผน", "ชั่วโมงที่ใช้"],
                [[p["name"], p["position"], p["total"], p["on_time"], p["late"],
                  p["missed"], p["cancelled"], p["still_open"],
                  "-" if p["on_time_pct"] is None else p["on_time_pct"] / 100,
                  round(p["minutes"] / 60, 1)] for p in people],
                [22, 18, 10, 10, 8, 10, 9, 11, 11, 13])
    for row in s2.iter_rows(min_row=2, min_col=9, max_col=9):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0%"

    # ── ชีตรายการ ──
    s3 = book.create_sheet("รายการทั้งหมด")
    write_sheet(s3, ["วันที่เริ่ม", "วันที่สิ้นสุด", "ชื่องาน", "ประเภท", "ผู้รับผิดชอบ",
                     "ผลลัพธ์", "ช้า (วัน)", "วันที่ทำเสร็จ", "นาทีที่ใช้", "บันทึกผล"],
                [[r["planned_date"], r["planned_end_date"], r["title"],
                  r["category_name"] or "", r["assignee_name"] or "",
                  reports.OUTCOME_LABELS.get(r["outcome"], r["outcome"]),
                  r["late_days"] or "", r["done_on"] or "",
                  r["actual_minutes"] if r["actual_minutes"] is not None else "",
                  r["result_note"] or r["cancel_reason"] or ""] for r in rows],
                [13, 13, 42, 20, 20, 16, 10, 13, 11, 40])
    for row in s3.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.number_format = "DD/MM/YYYY"
    for row in s3.iter_rows(min_row=2, min_col=8, max_col=8):
        for cell in row:
            cell.number_format = "DD/MM/YYYY"

    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)

    filename = f"rpf-calendar-{a.isoformat()}-{b.isoformat()}.xlsx"
    return StreamingResponse(buffer, media_type=XLSX_MIME, headers={
        "Content-Disposition": f'attachment; filename="{filename}"',
    })
